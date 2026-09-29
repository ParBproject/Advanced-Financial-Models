import unittest
from pathlib import Path

import openpyxl

from financial_models.cash_flow import DEFAULT_CAPEX, CashFlowAssumptions, forecast_cash_flow
from financial_models.portfolio import portfolio_metrics, workbook_balanced_portfolio


WORKBOOK = Path(__file__).resolve().parents[1] / "Advanced_Financial_Models.xlsx"


class WorkbookConsistencyTests(unittest.TestCase):
    @classmethod
    def setUpClass(cls):
        cls.formulas = openpyxl.load_workbook(WORKBOOK, data_only=False)
        cls.values = openpyxl.load_workbook(WORKBOOK, data_only=True)

    def test_cash_flow_identity_matches_the_python_forecast(self):
        sheet = self.formulas["Cash Flow Model"]
        assumptions = CashFlowAssumptions()

        self.assertEqual(sheet["B5"].value, assumptions.starting_cash)
        self.assertEqual(sheet["B6"].value, assumptions.revenue_month1)
        self.assertEqual(sheet["B7"].value, assumptions.monthly_revenue_growth)
        self.assertEqual(sheet["B8"].value, assumptions.operating_expense_ratio)
        self.assertEqual(sheet["B9"].value, assumptions.monthly_loan_payment)
        self.assertEqual(sheet["B10"].value, assumptions.annual_discount_rate)
        self.assertEqual(sheet["B22"].value, assumptions.monthly_rent)

        for month in range(1, 13):
            cell = sheet.cell(row=25, column=month + 1)
            column = cell.column_letter
            self.assertEqual(cell.value, DEFAULT_CAPEX.get(month, 0.0))
            self.assertEqual(
                sheet[f"{column}26"].value,
                f"={column}20+{column}22+{column}24+{column}25",
            )
            self.assertEqual(sheet[f"{column}27"].value, f"={column}18-{column}26")
            self.assertEqual(sheet[f"{column}28"].value, f"={column}15+{column}27")

        self.assertEqual(sheet["B32"].value, "=NPV((1+$B$10)^(1/12)-1,B27:M27)")

        cached = self.values["Cash Flow Model"]
        forecast = forecast_cash_flow()
        self.assertAlmostEqual(cached["B27"].value, 27_000.0)
        self.assertAlmostEqual(cached["B27"].value, forecast.periods[0].net_cash_flow)
        self.assertAlmostEqual(cached["M17"].value, forecast.periods[-1].revenue)
        self.assertAlmostEqual(cached["M28"].value, forecast.periods[-1].ending_cash)
        self.assertAlmostEqual(cached["B32"].value, forecast.npv_of_net_cash_flows)
        self.assertAlmostEqual(cached["B33"].value, forecast.average_monthly_net_cash_flow)
        self.assertAlmostEqual(cached["B34"].value, forecast.minimum_cash_balance)
        self.assertAlmostEqual(cached["B35"].value, forecast.maximum_cash_balance)

    def test_portfolio_risk_matches_zero_correlation_metrics(self):
        sheet = self.formulas["Portfolio Allocation"]
        assets = workbook_balanced_portfolio()

        for row, asset in enumerate(assets, start=11):
            self.assertEqual(sheet.cell(row=row, column=1).value, asset.name)
            self.assertEqual(sheet.cell(row=row, column=2).value, asset.expected_return)
            self.assertEqual(sheet.cell(row=row, column=3).value, asset.volatility)
            self.assertEqual(sheet.cell(row=row, column=4).value, asset.weight)

        self.assertEqual(
            sheet["B21"].value,
            "=SQRT(SUMPRODUCT(D11:D16,D11:D16,C11:C16,C11:C16))",
        )

        metrics = portfolio_metrics(assets)
        cached = self.values["Portfolio Allocation"]
        self.assertAlmostEqual(cached["B20"].value, metrics.expected_return)
        self.assertAlmostEqual(cached["B21"].value, metrics.volatility)
        self.assertAlmostEqual(cached["B22"].value, metrics.sharpe_ratio)
        self.assertAlmostEqual(cached["B23"].value, metrics.future_value)
        self.assertAlmostEqual(cached["B21"].value, 0.06824001758499187)

    def test_expected_loss_formula_is_exposure_times_pd_times_lgd(self):
        sheet = self.formulas["Loan Risk Model"]

        self.assertEqual(sheet["B5"].value, 0.45)
        self.assertEqual(sheet["E6"].value, 0.01)
        self.assertEqual(sheet["E10"].value, 0.15)
        self.assertEqual(sheet["E13"].value, 0.5)
        self.assertEqual(sheet["H17"].value.replace(" ", ""), "=G17*E17*F17")
        self.assertEqual(sheet["F6"].value, "Low")
        self.assertEqual(sheet["F10"].value, "Medium")
        self.assertEqual(sheet["F13"].value, "High")


if __name__ == "__main__":
    unittest.main()
